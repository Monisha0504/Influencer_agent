"""
FastAPI entry point for the Influencer Discovery Agent.

Run locally:
    uvicorn main:app --reload
Then open http://127.0.0.1:8000
"""

import io
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from fastapi import FastAPI, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agent import run_discovery, VALID_PLATFORMS

BASE_DIR = Path(__file__).resolve().parent

app = FastAPI(title="Influencer Discovery Agent")

app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")

# Cache-bust static/style.css so browsers pick up CSS changes after each deploy
# instead of serving a stale cached copy.
_STYLE_CSS_PATH = BASE_DIR / "static" / "style.css"
templates.env.globals["asset_version"] = (
    int(_STYLE_CSS_PATH.stat().st_mtime) if _STYLE_CSS_PATH.exists() else 0
)


# Curated dropdown lists shown to the user
CATEGORIES = [
    "Finance",
    "Food Blog",
    "Marketing",
    "Technology",
    "SaaS",
    "Fashion",
    "Beauty",
    "Fitness & Health",
    "Travel",
    "Lifestyle",
    "Gaming",
    "Education",
    "Parenting",
    "Automotive",
    "Real Estate",
    "Photography",
    "Music",
    "Comedy / Entertainment",
    "Business / Entrepreneurship",
    "Sustainability",
]

FOLLOWER_RANGES = [
    "Nano (1K - 10K)",
    "Micro (10K - 100K)",
    "Mid (100K - 500K)",
    "Macro (500K - 1M)",
    "Mega (1M+)",
]


# ---------- In-memory result cache (keyed by short token) ----------
# Lives until the server restarts. Good enough for a single-user local app.
_RESULTS_CACHE: dict[str, dict] = {}


def _store_result(result: dict, inputs: dict) -> str:
    token = uuid.uuid4().hex[:12]
    _RESULTS_CACHE[token] = {"result": result, "inputs": inputs}
    # Keep cache small - drop oldest if > 50 entries
    if len(_RESULTS_CACHE) > 50:
        for old in list(_RESULTS_CACHE.keys())[:-50]:
            _RESULTS_CACHE.pop(old, None)
    return token


# ---------- Routes ----------
@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse(
        request,
        "index.html",
        {
            "categories": CATEGORIES,
            "follower_ranges": FOLLOWER_RANGES,
            "platforms": VALID_PLATFORMS,
        },
    )


def _build_preference_directive(text: str) -> str:
    return (
        "CREATOR PREFERENCE — MANDATORY:\n"
        f"   - {text}\n"
        "   - Only include creators who clearly match this preference; do not include\n"
        "     creators just because they fit the category if they don't match this."
    )


@app.post("/discover", response_class=HTMLResponse)
async def discover(
    request: Request,
    category: str = Form(...),
    category_custom: Optional[str] = Form(None),
    followers_mode: str = Form("preset"),
    followers_range: Optional[str] = Form(None),
    followers_min: Optional[str] = Form(None),
    followers_max: Optional[str] = Form(None),
    website: Optional[str] = Form(None),
    location: Optional[str] = Form(None),
    platforms: List[str] = Form(default=[]),
    creator_preference: Optional[str] = Form(None),
):
    final_category = (category_custom or "").strip() or category

    if followers_mode == "manual" and (followers_min or followers_max):
        lo = (followers_min or "").strip() or "0"
        hi = (followers_max or "").strip() or "∞"
        final_followers = f"{lo} - {hi}"
    else:
        final_followers = followers_range or FOLLOWER_RANGES[1]

    final_website = (website or "").strip() or None
    final_location = (location or "").strip() or None
    final_platforms = [p for p in platforms if p in VALID_PLATFORMS] or None
    final_preference = (creator_preference or "").strip() or None
    preference_label = final_preference or "No preference"
    preference_directive = _build_preference_directive(final_preference) if final_preference else None

    token = None
    try:
        result = run_discovery(
            category=final_category,
            followers_range=final_followers,
            website=final_website,
            location=final_location,
            platforms=final_platforms,
            preference_directive=preference_directive,
        )
        error = None
        inputs = {
            "category": final_category,
            "followers_range": final_followers,
            "website": final_website or "",
            "location": final_location or "",
            "platforms": final_platforms or [],
            "platforms_display": ", ".join(final_platforms) if final_platforms else "All platforms",
            "creator_preference": final_preference or "",
            "creator_preference_label": preference_label,
        }
        token = _store_result(result, inputs)
    except Exception as e:
        result = None
        error = str(e)
        inputs = {
            "category": final_category,
            "followers_range": final_followers,
            "website": final_website or "",
            "location": final_location or "",
            "platforms": final_platforms or [],
            "platforms_display": ", ".join(final_platforms) if final_platforms else "All platforms",
            "creator_preference": final_preference or "",
            "creator_preference_label": preference_label,
        }

    return templates.TemplateResponse(
        request,
        "results.html",
        {
            "result": result,
            "error": error,
            "inputs": inputs,
            "download_token": token,
        },
    )


@app.get("/download/{token}.xlsx")
async def download_xlsx(token: str):
    entry = _RESULTS_CACHE.get(token)
    if not entry:
        raise HTTPException(status_code=404, detail="Result not found or expired. Please run the search again.")

    result = entry["result"]
    inputs = entry["inputs"]
    xlsx_bytes = _build_xlsx(result, inputs)

    safe_cat = "".join(c for c in inputs["category"] if c.isalnum() or c in ("_", "-")) or "category"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"influencers_{safe_cat}_{stamp}.xlsx"

    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/health")
async def health():
    return {"status": "ok"}


# ─── JSON API endpoint (used by ADMI Gateway / dashboard) ─────────────────────
from pydantic import BaseModel as _BM

class DiscoverRequest(_BM):
    category: str
    followers_range: str = "Micro (10K - 100K)"
    website: Optional[str] = None
    location: Optional[str] = None
    platforms: Optional[List[str]] = None
    creator_preference: Optional[str] = None

@app.post("/api/discover")
async def discover_api(payload: DiscoverRequest):
    """JSON version of /discover — returns structured influencer data."""
    final_preference = (payload.creator_preference or "").strip() or None
    try:
        result = run_discovery(
            category=payload.category,
            followers_range=payload.followers_range,
            website=payload.website,
            location=payload.location,
            platforms=payload.platforms,
            preference_directive=_build_preference_directive(final_preference) if final_preference else None,
        )
        return {
            "success": True,
            "category": payload.category,
            "followers_range": payload.followers_range,
            "total": len(result.get("influencers", [])),
            **result,
        }
    except Exception as e:
        return {"success": False, "error": str(e), "influencers": []}


# ---------- Excel builder ----------
def _build_xlsx(result: dict, inputs: dict) -> bytes:
    wb = Workbook()

    # ── Sheet 1: Summary ─────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"

    ws1.append(["Influencer Discovery Report"])
    ws1["A1"].font = Font(bold=True, size=18, color="4F46E5")
    ws1.merge_cells("A1:B1")
    ws1.row_dimensions[1].height = 32
    ws1.append([])

    meta = [
        ("Category",              inputs.get("category", "")),
        ("Followers Range",       inputs.get("followers_range", "")),
        ("Location Filter",       inputs.get("location", "") or "—"),
        ("Platform Filter",       inputs.get("platforms_display", "") or "All platforms"),
        ("Creator Preference",    inputs.get("creator_preference_label", "") or "No preference"),
        ("Website",               inputs.get("website", "") or "—"),
        ("Industry",              result.get("industry", "")),
        ("Target Audience",       result.get("target_audience", "")),
        ("Recommended Platforms", ", ".join(result.get("recommended_platforms", []) or [])),
        ("Total Influencers",     len(result.get("influencers", []) or [])),
        ("Generated At",          datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for label, val in meta:
        ws1.append([label, val])
        ws1.cell(row=ws1.max_row, column=1).font = Font(bold=True)
        ws1.cell(row=ws1.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 72

    # ── Sheet 2: Influencers ──────────────────────────────────────────
    ws2 = wb.create_sheet("Influencers")

    HDR_FONT  = Font(bold=True, color="FFFFFF", size=11)
    HDR_FILL  = PatternFill("solid", fgColor="4F46E5")
    ALT_FILL  = PatternFill("solid", fgColor="EEF2FF")
    GREEN     = PatternFill("solid", fgColor="D1FAE5")
    YELLOW    = PatternFill("solid", fgColor="FEF9C3")
    RED_FILL  = PatternFill("solid", fgColor="FEE2E2")
    LINK_FONT = Font(color="1D4ED8", underline="single", bold=True)
    THIN = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    headers = [
        "#", "Name", "Platform", "Niche", "Location",
        "Followers", "Engagement Rate", "Fit Score",
        "Fit Reason", "Preference Match", "Competitor Collab", "Content Strategy", "Profile Link",
    ]
    col_widths = [4, 24, 13, 22, 22, 15, 16, 11, 44, 44, 30, 44, 22]

    ws2.append(headers)
    ws2.row_dimensions[1].height = 32
    for i, cell in enumerate(ws2[1], start=1):
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = THIN
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws2.column_dimensions[get_column_letter(i)].width = col_widths[i - 1]

    for row_num, inf in enumerate(result.get("influencers", []) or [], start=2):
        fit = inf.get("fit_score", 0) or 0
        try:
            fit_int = int(fit)
        except Exception:
            fit_int = 0

        ws2.append([
            row_num - 1,
            inf.get("name", ""),
            inf.get("platform", ""),
            inf.get("niche", ""),
            inf.get("location", ""),
            inf.get("followers", ""),
            inf.get("engagement_rate", ""),
            fit_int,
            inf.get("fit_reason", ""),
            inf.get("preference_match", ""),
            inf.get("competitor_collaboration", ""),
            inf.get("content_strategy", ""),
            "View Profile",           # display text; hyperlink set below
        ])

        # Alternating row background
        row_fill = ALT_FILL if row_num % 2 == 0 else None

        for col_idx, cell in enumerate(ws2[row_num], start=1):
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_fill and col_idx != 8:
                cell.fill = row_fill

        # Colour-code Fit Score (column 8)
        score_cell = ws2.cell(row=row_num, column=8)
        score_cell.font = Font(bold=True)
        score_cell.alignment = Alignment(horizontal="center", vertical="top")
        if fit_int >= 80:
            score_cell.fill = GREEN
        elif fit_int >= 60:
            score_cell.fill = YELLOW
        else:
            score_cell.fill = RED_FILL

        # Clickable hyperlink in Profile Link column (column 13)
        url = inf.get("profile_link", "")
        if url and url.startswith("http"):
            link_cell = ws2.cell(row=row_num, column=13)
            link_cell.hyperlink = url
            link_cell.value = "View Profile"
            link_cell.font = LINK_FONT
            link_cell.alignment = Alignment(horizontal="center", vertical="top")

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
    